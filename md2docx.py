import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Customizing default styles
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_data = []
    
    for line in lines:
        line = line.strip()
        if not line:
            if in_table:
                # build table
                build_table(doc, table_data)
                in_table = False
                table_data = []
            continue
            
        if line.startswith('|'):
            in_table = True
            # extract cells
            cells = [cell.strip() for cell in line.split('|') if cell.strip() or len(cell.strip())==0][1:-1] if line.endswith('|') else [cell.strip() for cell in line.split('|')][1:]
            if all(set(c) <= {'-', ':', ' '} for c in cells):
                continue # skip separator row
            table_data.append(cells)
            continue
        
        if in_table:
            build_table(doc, table_data)
            in_table = False
            table_data = []
            
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('# ').replace('**', '')
            heading = doc.add_heading(text, level=min(level, 9))
            if level == 1:
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('* ') or line.startswith('- '):
            text = line[2:].replace('**', '')
            doc.add_paragraph(text, style='List Bullet')
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
            text = line[3:].replace('**', '')
            doc.add_paragraph(text, style='List Number')
        else:
            text = line.replace('**', '')
            doc.add_paragraph(text)

    if in_table:
        build_table(doc, table_data)

    doc.save(docx_path)

def build_table(doc, table_data):
    if not table_data: return
    cols = max(len(row) for row in table_data)
    table = doc.add_table(rows=len(table_data), cols=cols)
    table.style = 'Table Grid'
    for i, row in enumerate(table_data):
        for j, cell_text in enumerate(row):
            if j < cols:
                cell = table.cell(i, j)
                cell.text = cell_text.replace('**', '')

folder = 'docs/official_handover'
for f in os.listdir(folder):
    if f.endswith('.md'):
        md_path = os.path.join(folder, f)
        docx_path = os.path.join(folder, f.replace('.md', '.docx'))
        convert_md_to_docx(md_path, docx_path)
        print(f'Converted {f} to docx')
