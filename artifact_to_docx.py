import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_data = []
    
    for line in lines:
        line = line.strip()
        if not line:
            if in_table:
                build_table(doc, table_data)
                in_table = False
                table_data = []
            continue
            
        if line.startswith('|'):
            in_table = True
            raw_cells = [cell.strip() for cell in line.split('|')]
            if line.endswith('|'):
                cells = raw_cells[1:-1]
            else:
                cells = raw_cells[1:]
            
            if len(cells) > 0 and all(set(c) <= {'-', ':', ' '} for c in cells):
                continue
            table_data.append(cells)
            continue
        
        if in_table:
            build_table(doc, table_data)
            in_table = False
            table_data = []
            
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('# ').replace('**', '')
            heading = doc.add_heading(text, level=min(level, 9))
            if level == 1:
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('* ') or line.startswith('- '):
            text = line[2:].replace('**', '')
            doc.add_paragraph(text, style='List Bullet')
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
            text = line[3:].replace('**', '')
            doc.add_paragraph(text, style='List Number')
        else:
            text = line.replace('**', '')
            doc.add_paragraph(text)

    if in_table:
        build_table(doc, table_data)

    doc.save(docx_path)

def build_table(doc, table_data):
    if not table_data: return
    cols = max(len(row) for row in table_data)
    table = doc.add_table(rows=len(table_data), cols=cols)
    table.style = 'Table Grid'
    for i, row in enumerate(table_data):
        for j, cell_text in enumerate(row):
            if j < cols:
                cell = table.cell(i, j)
                cell.text = cell_text.replace('**', '')

if __name__ == '__main__':
    convert_md_to_docx(sys.argv[1], sys.argv[2])
