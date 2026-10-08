import sys
from docx import Document

def replace_in_paragraph(p, old_text, new_text):
    if old_text not in p.text:
        return
    
    # Best effort to keep run formatting if it's within a single run
    for run in p.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return

    # If it spans multiple runs, fallback to replacing paragraph text
    if old_text in p.text:
        new_full_text = p.text.replace(old_text, new_text)
        if len(p.runs) > 0:
            p.runs[0].text = new_full_text
            for i in range(1, len(p.runs)):
                p.runs[i].text = ""

def process_docx(input_path, output_path):
    doc = Document(input_path)
    
    # 1. REMOVE INTERNAL TECHNICAL REFERENCES
    
    p_replacements = {
        "Verification Evidence: Confirmed in `the report service` via atomic transaction rules and delta logic.": "Verification Evidence: Confirmed via atomic transaction rules and delta logic.",
        "Current Behaviour: Strict data siloing. Supervisors only see data for `farmIds` associated with their user profile. Verified via Firestore `rules` and React context filtering.": "Current Behaviour: Strict data siloing. Supervisors only see data for farms associated with their user profile."
    }
    
    for p in doc.paragraphs:
        for old_t, new_t in p_replacements.items():
            replace_in_paragraph(p, old_t, new_t)
            
    # Traceability table & Prediction
    t_replacements = {
        "`the prediction service`": "Prediction functionality",
        "`i18n` JSON files": "Language support",
        "`the report service` V2 logic": "V2 correction logic",
        "Blends WMA regression with standard curves": "Uses historical data and standard curves"
    }

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old_t, new_t in t_replacements.items():
                        replace_in_paragraph(p, old_t, new_t)

    # 2. CUSTOMER SIGN-OFF FIELDS
    # Add below existing sign-off fields
    for i, p in enumerate(doc.paragraphs):
        if "Remarks:" in p.text:
            # We want to insert AFTER this paragraph. 
            # In python-docx, insert_paragraph_before inserts BEFORE a paragraph. 
            # To insert AFTER, we can insert BEFORE the next paragraph, or if it's the last, just add_paragraph.
            if i + 1 < len(doc.paragraphs):
                next_p = doc.paragraphs[i+1]
                next_p.insert_paragraph_before("Customer Representative: ______________________")
                next_p.insert_paragraph_before("Designation: ______________________")
                next_p.insert_paragraph_before("Signature: ______________________")
            else:
                doc.add_paragraph("Customer Representative: ______________________")
                doc.add_paragraph("Designation: ______________________")
                doc.add_paragraph("Signature: ______________________")
            break

    doc.save(output_path)

if __name__ == '__main__':
    process_docx(sys.argv[1], sys.argv[2])
