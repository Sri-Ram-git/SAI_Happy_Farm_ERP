import sys
from docx import Document

def replace_text_in_paragraph(p, old_text, new_text):
    if old_text not in p.text:
        return
    
    for run in p.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return

    if old_text in p.text:
        new_full_text = p.text.replace(old_text, new_text)
        if len(p.runs) > 0:
            first_run = p.runs[0]
            first_run.text = new_full_text
            for i in range(1, len(p.runs)):
                p.runs[i].text = ""

def process_docx(input_path, output_path):
    doc = Document(input_path)
    
    replacements = {
        "Official / Final (Verified Implementation)": "Release 1.0 – Final Implementation Documentation",
        "Import/Revert: Import/Revert: Implemented & Verified": "Import/Revert: Implemented & Verified",
        "Verification Evidence: Confirmed in reportService.ts via atomic transaction rules and delta logic.": "Verification Evidence: Confirmed through atomic transaction and delta logic.",
        "db.runTransaction()": "atomic transactions",
        "collectionGroup('dailyLogs')": "aggregated data fetching",
        "reportService.ts": "the report service",
        "predictionService.ts": "the prediction service",
        "FarmerFormPage.tsx": "Farmer Portal",
        "kpiCalculations.ts": "KPI engine",
        "Firestore rules / React context filtering": "security rules and role-based filtering",
        "excelExportService.ts": "export service",
        "AdminImportPage.tsx": "Admin Import Portal",
        "i18n/hi.json": "localization configuration",
        "Weighted Moving Average and Linear Regression blended with Cobb/Ross standard curves": "historical data trends blended with standard breeder curves",
        "0 to 100 PPM": "0 to 50 PPM",
        "0-100 PPM": "0-50 PPM",
        "0–100 PPM": "0-50 PPM",
        "150": "60"
    }
    
    for p in doc.paragraphs:
        for old_t, new_t in replacements.items():
            replace_text_in_paragraph(p, old_t, new_t)
            
        if "EXECUTIVE SYSTEM OVERVIEW" in p.text:
            p.insert_paragraph_before("Release Scope Note: Release 1.0 implements the Local-Language Web Form approach (Option 2). The WhatsApp + OCR workflow is not included in this release.")

    for table in doc.tables:
        for row in table.rows:
            is_damaged_or_floor = False
            for cell in row.cells:
                if "Damaged Eggs" in cell.text or "Floor Eggs" in cell.text:
                    is_damaged_or_floor = True
                    break
            
            for cell in row.cells:
                if is_damaged_or_floor and "Enter saleable eggs" in cell.text:
                    if "Damaged Eggs" in row.cells[0].text or "Damaged Eggs" in row.cells[1].text:
                        replace_text_in_paragraph(cell.paragraphs[0], "Enter saleable eggs", "Enter damaged egg quantity")
                    elif "Floor Eggs" in row.cells[0].text or "Floor Eggs" in row.cells[1].text:
                        replace_text_in_paragraph(cell.paragraphs[0], "Enter saleable eggs", "Enter floor egg quantity")

                for cp in cell.paragraphs:
                    for old_t, new_t in replacements.items():
                        replace_text_in_paragraph(cp, old_t, new_t)

    found_signoff = False
    for p in doc.paragraphs:
        if "Customer Representative:" in p.text:
            if "Designation:" not in p.text:
                p.text = p.text + "\nDesignation: _______________________\nSignature: _______________________\nDate: _______________________"
            found_signoff = True
            break
            
    if not found_signoff:
        for p in doc.paragraphs:
            if "Tester Name:" in p.text:
                p.insert_paragraph_before("Customer Representative: _______________________")
                p.insert_paragraph_before("Designation: _______________________")
                p.insert_paragraph_before("Signature: _______________________")
                p.insert_paragraph_before("Date: _______________________")
                break

    doc.save(output_path)

if __name__ == '__main__':
    process_docx(sys.argv[1], sys.argv[2])
