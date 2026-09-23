import sys
import json
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference

# Standard Production Curves Data
CF_STD_CURVE = [
    (20, 0), (21, 2), (22, 8), (23, 22), (24, 42), (25, 62), (26, 76), (27, 83), (28, 85),
    (29, 85.5), (30, 85), (31, 84.5), (32, 84), (33, 83), (34, 82), (35, 81), (36, 80),
    (37, 79), (38, 78), (39, 77), (40, 76), (41, 74.5), (42, 73), (43, 71.5), (44, 70),
    (45, 68.5), (46, 67), (47, 65.5), (48, 64), (49, 62), (50, 60), (51, 58), (52, 56),
    (53, 54), (54, 52), (55, 50), (56, 48), (57, 46), (58, 44), (59, 42), (60, 40),
    (61, 38), (62, 36), (63, 34), (64, 32), (65, 30)
]

FR_STD_CURVE = [
    (20, 0), (21, 3), (22, 12), (23, 30), (24, 52), (25, 70), (26, 80), (27, 84), (28, 85),
    (29, 84.5), (30, 84), (31, 83), (32, 82), (33, 81), (34, 80), (35, 79), (36, 77.5),
    (37, 76), (38, 75), (39, 73.5), (40, 72), (41, 70.5), (42, 69), (43, 67.5), (44, 66),
    (45, 64.5), (46, 63), (47, 61.5), (48, 60), (49, 58), (50, 56), (51, 54), (52, 52),
    (53, 50), (54, 48), (55, 46), (56, 44), (57, 42), (58, 40), (59, 38), (60, 36),
    (61, 34), (62, 32), (63, 30), (64, 28), (65, 26)
]

def get_std_prod_pct(curve_type, age_weeks):
    curve = CF_STD_CURVE if curve_type == 'CF_STD' else FR_STD_CURVE
    if age_weeks <= curve[0][0]:
        return curve[0][1]
    if age_weeks >= curve[-1][0]:
        return curve[-1][1]
    
    for i in range(len(curve) - 1):
        w1, p1 = curve[i]
        w2, p2 = curve[i+1]
        if w1 <= age_weeks <= w2:
            frac = (age_weeks - w1) / (w2 - w1)
            return p1 + frac * (p2 - p1)
    return 0.0

def build_workbook(payload):
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    # 1. CF STD Sheet
    ws_cf = wb.create_sheet(title="CF STD")
    ws_cf.append(["Age (Weeks)", "Prod %"])
    for w, p in CF_STD_CURVE:
        ws_cf.append([w, p])

    # 2. FR STD Sheet
    ws_fr = wb.create_sheet(title="FR STD")
    ws_fr.append(["Age (Weeks)", "Prod %"])
    for w, p in FR_STD_CURVE:
        ws_fr.append([w, p])

    # Fonts & Styles
    font_calibri = Font(name="Calibri", size=11)
    font_bold = Font(name="Calibri", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    thin_border_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    farms_data = payload.get("farms", [])

    for farm in farms_data:
        farm_id = str(farm.get("farmId", "FARM")).strip()
        farm_name = str(farm.get("farmName", farm_id)).strip() # Worksheet name
        sheet_title = farm_name[:31] if farm_name else "FARM"
        
        ws = wb.create_sheet(title=sheet_title)

        # Label definitions in Column B starting Row 25
        labels = [
            (25, "Date"),
            (26, "WEEKS"),
            (27, "NO.OF BIRDS"),
            (28, "PRODUCTION"),
            (29, "SELECTION"),
            (30, "SELECTION %"),
            (31, "DAMAGE/REJECTED"),
            (32, "Mortality"),
            (33, "Temp"),
            (34, "Feed Kgs"),
            (35, "Feed Gms/Bird"),
            (36, "STD %"),
            (37, "ACT %"),
        ]

        for r_idx, label in labels:
            cell = ws.cell(row=r_idx, column=2, value=label)
            cell.font = font_bold if r_idx in (25, 26, 37) else font_calibri
            cell.alignment = align_left
            if r_idx == 37:
                cell.fill = yellow_fill

        reports = farm.get("reports", [])
        if not reports:
            continue

        # Sort chronologically by date
        reports.sort(key=lambda r: r.get("submissionDate", ""))

        curve_type = farm.get("productionCurve", "CF_STD")
        initial_birds = farm.get("initialBirds", 1000)

        start_col = 3 # Column C
        last_col = start_col + len(reports) - 1

        for idx, rep in enumerate(reports):
            col = start_col + idx
            col_letter = openpyxl.utils.get_column_letter(col)
            
            sub_date = rep.get("submissionDate", "")
            # Format date as MM-DD-YY if string YYYY-MM-DD
            if len(sub_date) == 10 and sub_date[4] == '-' and sub_date[7] == '-':
                parts = sub_date.split('-')
                formatted_date = f"{parts[1]}-{parts[2]}-{parts[0][2:]}"
            else:
                formatted_date = sub_date

            age_weeks = rep.get("ageWeeks", 17.0 + idx * 0.1)

            # Row 25: Date
            c25 = ws.cell(row=25, column=col, value=formatted_date)
            c25.font = font_calibri
            c25.alignment = align_center

            # Row 26: WEEKS
            c26 = ws.cell(row=26, column=col, value=round(age_weeks, 1))
            c26.font = font_bold
            c26.alignment = align_center
            c26.border = thin_border

            # Row 27: NO.OF BIRDS
            if idx == 0:
                c27 = ws.cell(row=27, column=col, value=rep.get("openingBirdCount", initial_birds))
            else:
                prev_col_letter = openpyxl.utils.get_column_letter(col - 1)
                c27 = ws.cell(row=27, column=col, value=f"={prev_col_letter}27-{prev_col_letter}32")
            c27.font = font_calibri
            c27.alignment = align_center
            c27.border = thin_border

            # Row 28: PRODUCTION
            prod_val = rep.get("eggsProduced", None)
            c28 = ws.cell(row=28, column=col, value=prod_val if prod_val is not None else "")
            c28.font = font_calibri
            c28.alignment = align_center
            c28.border = thin_border

            # Row 29: SELECTION
            sel_val = rep.get("selectionEggs", None)
            c29 = ws.cell(row=29, column=col, value=sel_val if sel_val is not None else "")
            c29.font = font_calibri
            c29.alignment = align_center
            c29.border = thin_border

            # Row 30: SELECTION %
            c30 = ws.cell(row=30, column=col, value=f"=IF(AND(ISNUMBER({col_letter}28), {col_letter}28>0, ISNUMBER({col_letter}29)), ({col_letter}29/{col_letter}28)*100, \"\")")
            c30.font = font_calibri
            c30.alignment = align_center
            c30.number_format = "0.00"
            c30.border = thin_border

            # Row 31: DAMAGE/REJECTED
            if prod_val is not None and sel_val is not None:
                dmg_val = max(0, prod_val - sel_val)
            else:
                dmg_val = ""
            c31 = ws.cell(row=31, column=col, value=dmg_val)
            c31.font = font_calibri
            c31.alignment = align_center
            c31.border = thin_border

            # Row 32: Mortality
            mort_val = rep.get("mortality", None)
            c32 = ws.cell(row=32, column=col, value=mort_val if (mort_val is not None and mort_val > 0) else "")
            c32.font = font_calibri
            c32.alignment = align_center
            c32.border = thin_border

            # Row 33: Temp
            temp_val = rep.get("temperature", None)
            c33 = ws.cell(row=33, column=col, value=temp_val if temp_val is not None else "")
            c33.font = font_calibri
            c33.alignment = align_center
            c33.border = thin_border

            # Row 34: Feed Kgs
            feed_val = rep.get("feedKg", None)
            c34 = ws.cell(row=34, column=col, value=feed_val if feed_val is not None else "")
            c34.font = font_calibri
            c34.alignment = align_center
            c34.border = thin_border

            # Row 35: Feed Gms/Bird
            c35 = ws.cell(row=35, column=col, value=f"=IF(AND(ISNUMBER({col_letter}27), {col_letter}27>0, ISNUMBER({col_letter}34)), ({col_letter}34*1000)/{col_letter}27, \"\")")
            c35.font = font_calibri
            c35.alignment = align_center
            c35.number_format = "0.00"
            c35.border = thin_border

            # Row 36: STD %
            std_pct = get_std_prod_pct(curve_type, age_weeks)
            c36 = ws.cell(row=36, column=col, value=round(std_pct, 2))
            c36.font = font_calibri
            c36.alignment = align_center
            c36.number_format = "0.00"
            c36.border = thin_border

            # Row 37: ACT %
            c37 = ws.cell(row=37, column=col, value=f"=IF(AND(ISNUMBER({col_letter}27), {col_letter}27>0, ISNUMBER({col_letter}28)), ({col_letter}28/{col_letter}27)*100, \"\")")
            c37.font = font_bold
            c37.alignment = align_center
            c37.fill = yellow_fill
            c37.number_format = "0.00"
            c37.border = thin_border

        # Adjust Column Widths
        ws.column_dimensions['A'].width = 3
        ws.column_dimensions['B'].width = 22
        for c in range(start_col, last_col + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 12

        # 1. Main Production Chart (Top at B2)
        main_chart = LineChart()
        main_chart.title = f"{sheet_title} - PRODUCTION"
        main_chart.style = 13
        main_chart.height = 10
        main_chart.width = 18

        xvalues = Reference(ws, min_col=start_col, min_row=26, max_col=last_col, max_row=26)
        act_series = Reference(ws, min_col=2, min_row=37, max_col=last_col, max_row=37)
        std_series = Reference(ws, min_col=2, min_row=36, max_col=last_col, max_row=36)

        main_chart.add_data(act_series, titles_from_data=True, from_rows=True)
        main_chart.add_data(std_series, titles_from_data=True, from_rows=True)
        main_chart.set_categories(xvalues)
        ws.add_chart(main_chart, "B2")

        # 2. Lower Chart 1: Selection Curve (B39)
        sel_chart = LineChart()
        sel_chart.title = "SELECTION CURVE"
        sel_chart.height = 7
        sel_chart.width = 11
        sel_series = Reference(ws, min_col=2, min_row=30, max_col=last_col, max_row=30)
        sel_chart.add_data(sel_series, titles_from_data=True, from_rows=True)
        sel_chart.set_categories(xvalues)
        ws.add_chart(sel_chart, "B39")

        # 3. Lower Chart 2: Feed Gms vs Production (I39)
        feed_chart = LineChart()
        feed_chart.title = "FEED GMS vs PRODUCTION"
        feed_chart.height = 7
        feed_chart.width = 11
        feed_series = Reference(ws, min_col=2, min_row=35, max_col=last_col, max_row=35)
        feed_chart.add_data(feed_series, titles_from_data=True, from_rows=True)
        feed_chart.add_data(act_series, titles_from_data=True, from_rows=True)
        feed_chart.set_categories(xvalues)
        ws.add_chart(feed_chart, "I39")

        # 4. Lower Chart 3: Production vs Temperature (P39)
        temp_chart = BarChart()
        temp_chart.title = "PRODUCTION vs TEMPERATURE"
        temp_chart.height = 7
        temp_chart.width = 11
        temp_series = Reference(ws, min_col=2, min_row=33, max_col=last_col, max_row=33)
        temp_chart.add_data(temp_series, titles_from_data=True, from_rows=True)
        temp_chart.add_data(act_series, titles_from_data=True, from_rows=True)
        temp_chart.set_categories(xvalues)
        ws.add_chart(temp_chart, "P39")

    # Safety Validation: Verify worksheets and references
    for ws_check in wb.worksheets:
        for chart in ws_check._charts:
            for s in chart.series:
                ref_str = s.val.numRef.f if (hasattr(s, 'val') and s.val and hasattr(s.val, 'numRef') and s.val.numRef) else str(getattr(s, 'val', s))
                # Ensure chart references ONLY its own worksheet or valid sheet title
                if "!" in ref_str:
                    sheet_in_ref = ref_str.split("!")[0].replace("'", "")
                    if sheet_in_ref != ws_check.title:
                        raise ValueError(f"Cross-farm reference detected in chart! Sheet '{ws_check.title}' references '{sheet_in_ref}'.")

    return wb

def main():
    if len(sys.argv) < 3:
        print("Usage: python generate_production_curve.py <input_json_path> <output_xlsx_path>", file=sys.stderr)
        sys.exit(1)

    input_json_path = sys.argv[1]
    output_xlsx_path = sys.argv[2]

    with open(input_json_path, 'r', encoding='utf-8') as f:
        payload = json.load(f)

    wb = build_workbook(payload)
    wb.save(output_xlsx_path)
    print(f"SUCCESS: {output_xlsx_path}")

if __name__ == "__main__":
    main()
